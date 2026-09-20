#!/usr/bin/env python3
"""Extract cached values of the legacy workbook used as regression expectations.

Produces ``backend/tests/fixtures/legacy_expected.json`` with, per article:
``initial_stock`` (stock at the first day of the grid), ``besoin`` and ``stock`` daily series
for the first ``--days`` days, and the weekly ``couverture`` / ``target`` values.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--out", default="backend/tests/fixtures/legacy_expected.json")
    ap.add_argument("--days", type=int, default=120)
    args = ap.parse_args()
    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb["SIMULATION"]
    dates = {c: ws.cell(7, c).value.date() for c in range(10, ws.max_column + 1)
             if isinstance(ws.cell(7, c).value, dt.datetime)}
    cols = sorted(dates)[: args.days]
    out: dict = {"start_date": dates[cols[0]].isoformat(), "dates": [dates[c].isoformat() for c in cols], "articles": {}}
    for r in range(9, ws.max_row + 1):
        ref, var = ws.cell(r, 3).value, ws.cell(r, 9).value
        if not ref or not var:
            continue
        m = re.fullmatch(r"(\d)\. ([^()]+?)(?: \((S-\d+)\))?", var)
        kind = m.group(2).strip()
        art = out["articles"].setdefault(ref, {})
        if kind == "Besoin":
            art["besoin"] = [float(ws.cell(r, c).value or 0) for c in cols]
        elif kind == "Stock":
            art["initial_stock"] = float(ws.cell(r, cols[0]).value)
            art["stock"] = [float(ws.cell(r, c).value) if isinstance(ws.cell(r, c).value, (int, float)) else None for c in cols]
        elif kind == "Couverture":
            art["couverture"] = {dates[c].isoformat(): int(ws.cell(r, c).value) for c in cols if isinstance(ws.cell(r, c).value, (int, float))}
        elif kind == "Target stock":
            art["target"] = {dates[c].isoformat(): float(ws.cell(r, c).value) for c in cols if isinstance(ws.cell(r, c).value, (int, float))}
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(out, indent=1))
    print(f"wrote {args.out}: {len(out['articles'])} articles, {len(cols)} days")


if __name__ == "__main__":
    main()
