"""Regression tests of the audit's corrections (report of 29/10/2026) : each test names the finding it covers."""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook, load_workbook
from test_api import AID, AS_OF, client  # noqa: F401
from test_engine_units import MON, make_dataset, run

from appro.engine.models import DesadvLine, Receipt, Supplier, SupplierLink
from appro.services import data_service


def _planners(c):
    """A root admin, quentin as an appro of his own portfolio, alice as another appro, a reader."""
    arts = c.get("/api/reference/ref_articles/rows").json()
    mine = next(a for a in arts if a["article_id"] == AID)["planner"]
    other = next(a for a in arts if a["article_id"] != AID)
    other["planner"] = "ALICE"
    assert c.put("/api/reference/ref_articles/rows", json={"values": other}).status_code == 200
    root = {"x-forwarded-email": "root@example.com"}
    put = lambda vals, h: c.put("/api/reference/ref_planners/rows", json={"values": vals}, headers=h)  # noqa: E731
    assert put({"planner_id": "PROC0", "name": "ROOT", "email": "root@example.com", "role": "admin", "active": True}, {}).status_code == 200
    assert put({"planner_id": "PROC1", "name": mine, "email": "quentin@example.com", "role": "appro", "active": True}, root).status_code == 200
    assert put({"planner_id": "PROC2", "name": "ALICE", "email": "alice@example.com", "role": "appro", "active": True}, root).status_code == 200
    return root, other["article_id"]


def test_f02_delete_override_needs_the_same_right_as_put(client):  # noqa: F811
    root, _other = _planners(client)
    reader = {"x-forwarded-email": "nobody@example.com"}
    r = client.put("/api/params/overrides", json={"scope": "global", "field": "target_tolerance_pct", "value": 50}, headers=root)
    assert r.status_code == 200
    oid = r.json()["id"]
    assert client.put("/api/params/overrides", json={"scope": "global", "field": "target_tolerance_pct", "value": 40}, headers=reader).status_code == 403
    assert client.delete(f"/api/params/overrides/{oid}", headers=reader).status_code == 403          # was 204
    assert client.delete(f"/api/params/overrides/{oid}").status_code == 403                          # an appro neither
    assert client.delete(f"/api/params/overrides/{oid}", headers=root).status_code == 204


def test_f03_undo_requires_current_rights(client):  # noqa: F811
    root, _other = _planners(client)
    day = (dt.date.fromisoformat(AS_OF) + dt.timedelta(days=3)).isoformat()
    r = client.put("/api/entries/adjustments", json={"article_id": AID, "date": day, "expression": "-5"})
    assert r.status_code == 200
    action = r.headers["X-Action-Id"]
    # quentin's declaration is deactivated : he is a reader now, his own action can no longer be undone
    assert client.put("/api/reference/ref_planners/rows", json={"values": {"planner_id": "PROC1", "name": "QUENTIN", "email": "quentin@example.com", "role": "appro", "active": False}}, headers=root).status_code == 200
    assert client.get("/api/config").json()["access"]["role"] == "reader"
    assert client.post(f"/api/actions/{action}/undo").status_code == 403
    assert client.post("/api/actions/undo-last").status_code == 403
    assert [a for a in client.get("/api/entries/adjustments", params={"article_id": AID}).json() if a["date"] == day]
    # an administrator keeps the reprise
    assert client.post(f"/api/actions/{action}/undo", headers=root).status_code == 200


def test_f03_undo_outside_the_portfolio_is_refused(client):  # noqa: F811
    root, other = _planners(client)
    day = (dt.date.fromisoformat(AS_OF) + dt.timedelta(days=3)).isoformat()
    # alice writes on her article, then her article moves to quentin's portfolio : alice can no longer revert it
    alice = {"x-forwarded-email": "alice@example.com"}
    r = client.put("/api/entries/adjustments", json={"article_id": other, "date": day, "expression": "-7"}, headers=alice)
    assert r.status_code == 200
    row = next(a for a in client.get("/api/reference/ref_articles/rows").json() if a["article_id"] == other)
    assert client.put("/api/reference/ref_articles/rows", json={"values": {**row, "planner": "QUENTIN"}}, headers=root).status_code == 200
    r = client.post(f"/api/actions/{r.headers['X-Action-Id']}/undo", headers=alice)
    assert r.status_code == 403 and "hors de votre carnet" in r.json()["detail"]


def test_f24_portal_export_is_allowed_to_a_supplier(client):  # noqa: F811
    rows = client.get("/api/reference/ref_suppliers/rows").json()
    r = next(x for x in rows if x["supplier_id"] == "S-000378")
    keep = {k: v for k, v in r.items() if k in ("supplier_id", "name", "country", "contact", "delivery_weekdays", "strict_days", "lead_time_days", "active", "docks")}
    assert client.put("/api/reference/ref_suppliers/rows", json={"values": {**keep, "transport_contact": "quai@fournisseur.example"}}).status_code == 200
    h = {"x-forwarded-email": "quai@fournisseur.example"}
    body = {"title": "Commandes", "sheet": "Commandes", "columns": [{"key": "a", "label": "A"}], "rows": [{"a": 1}], "footer": None}
    assert client.post("/api/exports/table.xlsx", json=body, headers=h).status_code == 200
    assert client.get("/api/cockpit", headers=h).status_code == 403


def test_f01_static_route_never_leaves_the_build(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from appro.api.main import create_app
    from appro.config import get_settings
    static = tmp_path / "dist"
    (static / "assets").mkdir(parents=True)
    (static / "index.html").write_text("<html>app</html>")
    (static / "brand.svg").write_text("<svg/>")
    (tmp_path / "secret.txt").write_text("sentinel")
    monkeypatch.setenv("APPRO_STATIC_DIR", str(static))
    get_settings.cache_clear()
    try:
        app = TestClient(create_app(), headers={"x-forwarded-email": "quentin@example.com"})
        assert app.get("/brand.svg").text == "<svg/>"
        for path in ("/../secret.txt", "/%2e%2e/secret.txt", "/assets/../../secret.txt", "/%2e%2e%2fsecret.txt"):
            r = app.get(path)
            assert "sentinel" not in r.text, path
    finally:
        get_settings.cache_clear()


def _pdp_wide(rows: list[list], header: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SOP - PDP"
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_f10_pdp_import_refuses_duplicates_negatives_and_uses_the_identifier(client):  # noqa: F811
    progs = client.get("/api/reference/ref_programs/rows").json()
    p1, p2 = progs[0], progs[1]
    header = ["PF/SF (nom du programme)", "program_id (facultatif)", "2026-W40", "2026-W41"]
    post = lambda content: client.post("/api/pdp/import", files={"file": ("pdp.xlsx", content)}, data={"name": "t", "activate": "false"})  # noqa: E731
    r = post(_pdp_wide([[p1["name"], p1["program_id"], 100, 200], [p1["name"], p1["program_id"], 300, 0]], header))
    assert r.status_code == 422 and "doublon" in r.json()["detail"]
    r = post(_pdp_wide([[p1["name"], p1["program_id"], 100, -5]], header))
    assert r.status_code == 422 and "négative" in r.json()["detail"]
    r = post(_pdp_wide([[p1["name"], p1["program_id"], "abc", 1]], header))
    assert r.status_code == 422 and "non numérique" in r.json()["detail"]
    # two programmes with the same display name : the identifier column decides
    r = post(_pdp_wide([[p1["name"], p1["program_id"], 100, 200], [p1["name"], p2["program_id"], 300, 400]], header))
    assert r.status_code == 201, r.text
    sheet = client.get("/api/pdp/sheet").json()
    vals = {p["program_id"]: p for p in sheet["programs"]}
    assert vals[p1["program_id"]]["source"] == "app" and vals[p2["program_id"]]["source"] == "app"
    # long layout : the same (programme, week) twice is refused, the decimal comma is accepted
    wb = Workbook()
    ws = wb.active
    ws.append(["program_id", "week_start", "qty"])
    ws.append([p1["program_id"], "2026-W42", "12,5"])
    ws.append([p1["program_id"], "2026-W42", 3])
    buf = io.BytesIO()
    wb.save(buf)
    r = post(buf.getvalue())
    assert r.status_code == 422 and "doublon" in r.json()["detail"]


def test_f11_a_real_first_row_is_not_an_example(client):  # noqa: F811
    t = client.get("/api/reference/ref_suppliers/template.xlsx")
    wb = load_workbook(io.BytesIO(t.content))
    ws = wb["Fournisseurs"]
    # the template's example key (S-000545) with another designation : a real row, kept
    ws["A2"], ws["B2"] = "S-000545", "ELANTAS renommé"
    buf = io.BytesIO()
    wb.save(buf)
    rep = client.post("/api/reference/ref_suppliers/import", files={"file": ("f.xlsx", buf.getvalue())}, data={"mode": "merge"})
    assert rep.status_code == 201, rep.text
    assert "ligne d'exemple ignorée" not in " ".join(rep.json().get("notes", []))
    assert next(r for r in client.get("/api/reference/ref_suppliers/rows").json() if r["supplier_id"] == "S-000545")["name"] == "ELANTAS renommé"
    # the untouched example row is still skipped
    t = client.get("/api/reference/ref_suppliers/template.xlsx")
    rep = client.post("/api/reference/ref_suppliers/import", files={"file": ("f.xlsx", t.content)}, data={"mode": "merge"})
    assert rep.status_code == 201 and "ligne d'exemple ignorée" in " ".join(rep.json().get("notes", []))


def test_f12_invalid_numbers_are_refused_and_the_decimal_comma_accepted(client):  # noqa: F811
    rows = client.get("/api/reference/ref_article_suppliers/rows").json()
    link = next(r for r in rows if r["article_id"] == AID)
    base = {k: v for k, v in link.items() if k not in ("updated_by", "updated_at")}
    r = client.put("/api/reference/ref_article_suppliers/rows", json={"values": {**base, "moq": "12,5"}})
    assert r.status_code == 200, r.text
    assert next(x for x in client.get("/api/reference/ref_article_suppliers/rows").json() if x["article_id"] == AID and x["supplier_id"] == link["supplier_id"])["moq"] == 12.5
    r = client.put("/api/reference/ref_article_suppliers/rows", json={"values": {**base, "moq": "pas un nombre"}})
    assert r.status_code == 422 and "n'est pas un nombre" in r.json()["detail"]
    r = client.put("/api/reference/ref_article_suppliers/rows", json={"values": {**base, "max_qty_per_week": "abc"}})
    assert r.status_code == 422


def test_f14_weekly_cap_below_moq_is_refused(client):  # noqa: F811
    rows = client.get("/api/reference/ref_article_suppliers/rows").json()
    link = next(r for r in rows if r["article_id"] == AID)
    base = {k: v for k, v in link.items() if k not in ("updated_by", "updated_at")}
    r = client.put("/api/reference/ref_article_suppliers/rows", json={"values": {**base, "moq": 100, "pack_qty": 50, "max_qty_per_week": 50, "max_qty_per_delivery": 100}})
    assert r.status_code == 422 and "≥ MOQ" in r.json()["detail"]


def test_f15_inactive_supplier_gets_no_proposal_but_keeps_its_history():
    ds = make_dataset(suppliers=[Supplier("S1", "Supplier 1", active=False)],
                      receipts=[Receipt("R1", "A1", MON - dt.timedelta(days=3), 200, supplier_id="S1")],
                      stock=[], )
    from appro.engine.models import StockSnapshot
    ds.stock = [StockSnapshot("A1", MON - dt.timedelta(days=8), 100.0)]
    r = run(ds, horizon_days=30)
    assert not [p for p in r.proposals if p.supplier_id == "S1"]      # no new proposal addressed to the inactive supplier
    lane = next(ln for ln in r.lanes if ln.supplier_id == "S1")
    assert float(lane.receipts.sum()) == 200.0           # the past receipt still counts
    ds2 = make_dataset()
    assert [p for p in run(ds2, horizon_days=30).proposals if p.supplier_id == "S1"]   # active : it does


def test_f17_desadv_matched_per_supplier_and_bl():
    mon1 = MON + dt.timedelta(days=7)
    ds = make_dataset(suppliers=[Supplier("S1", "Supplier 1"), Supplier("S2", "Supplier 2")],
                      links=[SupplierLink("A1", "S1", moq=100, pack_qty=50, quota_pct=50, priority=1), SupplierLink("A1", "S2", moq=100, pack_qty=50, quota_pct=50, priority=2)],
                      receipts=[Receipt("R1", "A1", MON - dt.timedelta(days=2), 100, supplier_id="S1", packing_slip="BL42", status="Reçu")])
    ds.desadv = [DesadvLine("D1", "A1", "S1", "BL42", MON - dt.timedelta(days=3), 100, state="Traité", final_processing="OK"),
                 DesadvLine("D2", "A1", "S2", "BL42", MON - dt.timedelta(days=3), 100, state="Traité", final_processing="OK")]
    r = run(ds, no_cbn=True, horizon_days=10)
    by = {(d.supplier_id if hasattr(d, "supplier_id") else ln.supplier_id, d.desadv_id): d for ln in r.lanes for d in ln.desadv}
    s1 = next(d for (sid, did), d in by.items() if did == "D1")
    s2 = next(d for (sid, did), d in by.items() if did == "D2")
    assert s1.received is True and s2.received is False     # the same BL number at another supplier is not received
    assert mon1 > MON


def test_f23_sync_status_flags_an_expected_table_never_copied(client):  # noqa: F811
    from appro.data.store import ErpSyncLog
    from appro.services.context import get_context
    ctx = get_context()
    original = ctx.settings
    try:
        ctx.settings = original.model_copy(update={"data_source": "lakebase"})
        with ctx.session() as session:
            session.add(ErpSyncLog(table_name="erp_purchase_orders", row_count=10, synced_at=dt.datetime.utcnow(), source="uc", run_id="1"))
            session.commit()
            st = data_service.sync_status(ctx, session)
        by = {t["table"]: t["status"] for t in st["tables"]}
        assert by["erp_purchase_orders"] == "ok" and by["erp_receipts"] == "missing"
        assert st["status"] == "late"                       # one fresh table does not make the whole copy fresh
        ctx.settings = original.model_copy(update={"data_source": "lakebase", "erp_desadv_table": None, "erp_prices_table": None,
                                                   "erp_stock_table": None, "erp_consumption_table": None})
        with ctx.session() as session:
            for t in ("erp_receipts",):
                session.add(ErpSyncLog(table_name=t, row_count=1, synced_at=dt.datetime.utcnow(), source="uc", run_id="1"))
            session.commit()
            st = data_service.sync_status(ctx, session)
        by = {t["table"]: t["status"] for t in st["tables"]}
        assert by["erp_desadv"] == "disabled" and st["status"] == "ok"
    finally:
        ctx.settings = original


def test_f34_health_reports_a_degraded_database_with_503(client):  # noqa: F811
    from appro.services.context import get_context
    ctx = get_context()
    factory = ctx.session_factory

    class Broken:
        def __call__(self, *a, **k):
            raise RuntimeError("postgresql://user:secret@host/db refused")

    ctx.session_factory = Broken()
    try:
        r = client.get("/api/health")
        assert r.status_code == 503
        body = r.json()
        assert body["status"] == "degraded" and body["database_status"] == "RuntimeError" and "secret" not in r.text
    finally:
        ctx.session_factory = factory
    assert client.get("/api/health").status_code == 200
