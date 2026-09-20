"""End-to-end API tests on an isolated context (seed CSV source + in-memory SQLite store)."""
from __future__ import annotations

import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.pool import StaticPool
from tests.conftest import SEED

from appro.api.main import create_app
from appro.config import Settings
from appro.data.store import Base
from appro.services.context import AppContext, set_context

AS_OF = "2026-09-19"


@pytest.fixture()
def client(seed_source):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    settings = Settings(data_source="local", seed_dir=SEED, as_of=dt.date.fromisoformat(AS_OF), horizon_days=90,
                        static_dir=SEED / "no-such-dir")
    ctx = AppContext(settings=settings, source=seed_source, session_factory=sessionmaker(bind=engine, expire_on_commit=False))
    set_context(ctx)
    try:
        yield TestClient(create_app(), headers={"x-forwarded-email": "quentin@example.com"})
    finally:
        set_context(None)


def test_health_and_config(client):
    assert client.get("/api/health").json()["status"] == "ok"
    cfg = client.get("/api/config").json()
    assert cfg["as_of"] == AS_OF and cfg["planners"] == ["QUENTIN"] and cfg["user"] == "quentin@example.com"


def test_cockpit(client):
    r = client.get("/api/cockpit", params={"planner": "QUENTIN"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["kpis"]["articles"] == 16
    assert len(body["articles"]) == 16
    assert body["as_of"] == AS_OF
    first = body["articles"][0]
    assert {"article_id", "kpis", "sparkline", "severity"} <= set(first)
    assert body["weekly_supply_demand"] and body["weekly_supply_demand"][0]["week"].startswith("2026-W")
    # proposals are not computed on display: some articles run short until the CBN is run
    assert body["kpis"]["sim_order_articles"] == 0
    for a in body["articles"]:
        assert a["kpis"]["min_stock_sim"] >= 0 and a["kpis"]["proposal_count"] == 0, a["article_id"]


def test_projection_day_and_week(client):
    r = client.get("/api/articles/P-00001046/projection", params={"granularity": "day"})
    assert r.status_code == 200, r.text
    day = r.json()
    keys = {s["key"] for s in day["series"]}
    assert {"demand", "stock_firm", "stock_sim", "coverage_sim", "target_stock"} <= keys
    assert len(day["periods"]) == len(day["series"][0]["values"])
    assert day["programs"] and day["suppliers"][0]["supplier_id"] == "S-000545"
    week = client.get("/api/articles/P-00001046/projection", params={"granularity": "week"}).json()
    assert len(week["periods"]) < len(day["periods"])
    demand_day = sum(next(s for s in day["series"] if s["key"] == "demand")["values"])
    demand_week = sum(next(s for s in week["series"] if s["key"] == "demand")["values"])
    assert abs(demand_day - demand_week) < 0.01
    assert client.get("/api/articles/NOPE/projection").status_code == 404


def test_entries_change_projection_and_audit(client):
    before = client.get("/api/articles/P-00001046/projection", params={"generate_proposals": "false"}).json()
    stock_before = next(s for s in before["series"] if s["key"] == "stock_sim")["values"]
    # an order typed in the app is a real (firm) order: simulated orders are grid cells
    assert client.post("/api/entries/orders", json={"article_id": "P-00001046", "supplier_id": "S-000545",
                                                    "expected_date": "2026-09-25", "qty": 1600, "order_type": "PLANNED"}).status_code == 422
    r = client.post("/api/entries/orders", json={"article_id": "P-00001046", "supplier_id": "S-000545",
                                                 "expected_date": "2026-09-25", "qty": 1600})
    assert r.status_code == 201, r.text
    order = r.json()
    after = client.get("/api/articles/P-00001046/projection", params={"generate_proposals": "false"}).json()
    stock_after = next(s for s in after["series"] if s["key"] == "stock_sim")["values"]
    i = after["periods"].index("2026-09-25")
    assert stock_after[i] - stock_before[i] == pytest.approx(1600)
    firm_after = next(s for s in after["series"] if s["key"] == "stock_firm")["values"]
    firm_before = next(s for s in before["series"] if s["key"] == "stock_firm")["values"]
    assert firm_after[i] - firm_before[i] == pytest.approx(1600)  # firm layer too
    # receipt closes the order
    r = client.post("/api/entries/receipts", json={"article_id": "P-00001046", "order_id": order["id"],
                                                   "receipt_date": "2026-09-24", "qty": 1600})
    assert r.status_code == 201
    assert client.get("/api/entries/orders", params={"article_id": "P-00001046"}).json()[0]["status"] == "RECEIVED"
    # adjustment + production + deletes
    assert client.post("/api/entries/adjustments", json={"article_id": "P-00001046", "date": "2026-09-22", "qty": -50}).status_code == 201
    assert client.post("/api/entries/adjustments", json={"article_id": "P-00001046", "date": "2026-09-22", "qty": 0}).status_code == 422
    assert client.post("/api/entries/production", json={"program_id": "mass-00040633", "date": "2026-09-21", "qty": 0}).status_code == 201
    assert client.post("/api/entries/production", json={"program_id": "nope", "date": "2026-09-21", "qty": 1}).status_code == 404
    log = client.get("/api/audit").json()
    assert log and log[0]["user"] == "quentin@example.com"
    actions = {e["action"] for e in log}
    assert {"create", "upsert"} <= actions


def test_cells_and_cbn_run(client):
    # type a simulated order with an expression, then an adjustment, on P-00001046
    r = client.put("/api/entries/cells", json={"article_id": "P-00001046", "date": "2026-09-25", "kind": "sim_order",
                                               "expression": "1000 + 2*100"})
    assert r.status_code == 200, r.text
    cell = r.json()
    assert cell["qty"] == 1200 and cell["expression"] == "1000 + 2*100" and cell["source"] == "MANUAL"
    assert client.put("/api/entries/cells", json={"article_id": "P-00001046", "date": "2026-09-25", "kind": "sim_order",
                                                  "expression": "abc"}).status_code == 422
    assert client.put("/api/entries/cells", json={"article_id": "NOPE", "date": "2026-09-25", "expression": "1"}).status_code == 404
    r = client.put("/api/entries/cells", json={"article_id": "P-00001046", "date": "2026-09-23", "kind": "adjustment",
                                               "expression": "-(30+20)"})
    assert r.json()["qty"] == -50
    proj = client.get("/api/articles/P-00001046/projection", params={"granularity": "day"}).json()
    series = {x["key"]: x["values"] for x in proj["series"]}
    i = proj["periods"].index("2026-09-25")
    assert series["supply_planned"][i] == 1200 and series["adjustments"][proj["periods"].index("2026-09-23")] == -50
    assert series["stock_sim"][i] - series["stock_forecast"][i] == pytest.approx(1200)
    assert "supply_proposed" not in series and "target_stock" in series
    # a negative simulated order is allowed ; an empty expression clears the cell
    assert client.put("/api/entries/cells", json={"article_id": "P-00001046", "date": "2026-09-26", "expression": "-300"}).json()["qty"] == -300
    assert client.put("/api/entries/cells", json={"article_id": "P-00001046", "date": "2026-09-26", "expression": ""}).json() is None
    assert len(client.get("/api/entries/cells", params={"article_id": "P-00001046"}).json()) == 2

    # a typed cell on the very day of a proposal: the proposal is added to it, the cell stays manual
    assert client.put("/api/entries/cells", json={"article_id": "P-00003751", "date": "2026-09-21", "expression": "2*300-100"}).json()["qty"] == 500
    # CBN run on the planner perimeter: proposals become simulated-order cells, no stockout remains
    r = client.post("/api/cbn/run", json={"planner": "QUENTIN"})
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["articles"] == 16 and rep["proposals"] > 0 and rep["qty"] > 0 and rep["removed"] == 0
    assert rep["items"][0]["proposal_id"].startswith("PR-")
    cells = client.get("/api/entries/cells", params={"kind": "sim_order"}).json()
    assert sum(1 for c in cells if c["source"] == "CBN") == rep["proposals"]
    manual = next(c for c in cells if c["article_id"] == "P-00001046" and c["date"] == "2026-09-25")
    assert manual["source"] == "MANUAL" and manual["qty"] >= 1200  # typed cells are kept
    same_day = [c for c in cells if c["article_id"] == "P-00003751" and c["date"] == "2026-09-21"]
    assert {c["source"] for c in same_day} == {"MANUAL", "CBN"} and next(c for c in same_day if c["source"] == "MANUAL")["qty"] == 500
    # typing on that day replaces both cells by the typed value
    assert client.put("/api/entries/cells", json={"article_id": "P-00003751", "date": "2026-09-21", "expression": "700"}).json()["qty"] == 700
    assert [c["source"] for c in client.get("/api/entries/cells", params={"article_id": "P-00003751"}).json() if c["date"] == "2026-09-21"] == ["MANUAL"]
    client.put("/api/entries/cells", json={"article_id": "P-00003751", "date": "2026-09-21", "expression": "2*300-100"})
    body = client.get("/api/cockpit", params={"planner": "QUENTIN"}).json()
    assert body["kpis"]["sim_order_articles"] > 0 and body["kpis"]["sim_orders_qty"] > 0
    for a in body["articles"]:
        assert a["kpis"]["first_stockout_sim"] is None, a["article_id"]
    # a second run is reproducible: previous CBN cells are replaced, typed cells respected
    rep2 = client.post("/api/cbn/run", json={"planner": "QUENTIN"}).json()
    assert rep2["removed"] > 0 and rep2["proposals"] == rep["proposals"]
    assert client.delete(f"/api/entries/cells/{manual['id']}").status_code == 204
    assert client.delete(f"/api/entries/cells/{manual['id']}").status_code == 404
    assert client.post("/api/cbn/run", json={"planner": "QUENTIN", "scenario_id": "nope"}).status_code == 404


def test_scenarios_and_simulate(client):
    body = {"name": "PDP +20 %", "description": "hausse", "params": {"horizon_days": 60},
            "events": [{"kind": "plan_factor", "payload": {"factor": 1.2}, "label": "+20 % partout"},
                       {"kind": "add_order", "payload": {"article_id": "P-00003751", "supplier_id": "S-000599",
                                                         "date": "2026-10-05", "qty": 2700}}]}
    r = client.post("/api/scenarios", json=body)
    assert r.status_code == 201, r.text
    sc = r.json()
    assert len(sc["events"]) == 2 and sc["params"]["horizon_days"] == 60
    cmp_ = client.get(f"/api/scenarios/{sc['id']}/compare", params={"planner": "QUENTIN"}).json()
    assert cmp_["scenario_kpis"]["demand_next_30d"] > cmp_["base_kpis"]["demand_next_30d"]
    assert len(cmp_["articles"]) == 16
    cockpit = client.get("/api/cockpit", params={"planner": "QUENTIN", "scenario_id": sc["id"]}).json()
    assert cockpit["horizon_days"] == 60
    assert client.get("/api/cockpit", params={"scenario_id": "nope"}).status_code == 404
    assert client.post("/api/scenarios", json={"name": "x", "events": [{"kind": "bogus"}]}).status_code == 422
    sim = client.post("/api/simulate", json={"planner": "QUENTIN", "article_ids": ["P-00001046"],
                                             "events": [{"kind": "move_order", "payload": {"order_id": "PO-000032", "days": 10}}],
                                             "params": {"generate_proposals": False}})
    assert sim.status_code == 200, sim.text
    assert len(sim.json()["articles"]) == 1
    r = client.put(f"/api/scenarios/{sc['id']}", json={**body, "name": "renommé", "events": body["events"][:1]})
    assert r.json()["name"] == "renommé" and len(r.json()["events"]) == 1
    assert client.delete(f"/api/scenarios/{sc['id']}").status_code == 204
    assert client.get(f"/api/scenarios/{sc['id']}").status_code == 404


def test_params_overrides(client):
    schema = client.get("/api/params/schema").json()
    assert any(p["field"] == "coverage_unit" and p["options"] for p in schema)
    r = client.put("/api/params/overrides", json={"scope": "global", "field": "horizon_days", "value": 45})
    assert r.status_code == 200, r.text
    assert client.get("/api/params/effective").json()["horizon_days"] == 45
    assert client.get("/api/cockpit").json()["horizon_days"] == 45
    r = client.put("/api/params/overrides", json={"scope": "article", "key1": "P-00001046", "field": "coverage_target_days", "value": 20})
    assert r.status_code == 200
    arts = client.get("/api/reference/articles").json()
    assert next(a for a in arts if a["article_id"] == "P-00001046")["coverage_target_days"] == 20
    assert client.put("/api/params/overrides", json={"scope": "global", "field": "not_a_field", "value": 1}).status_code == 422
    assert client.put("/api/params/overrides", json={"scope": "global", "field": "horizon_days", "value": "abc"}).status_code == 422
    rows = client.get("/api/params/overrides").json()
    assert len(rows) == 2
    assert client.delete(f"/api/params/overrides/{rows[0]['id']}").status_code == 204


def test_reference_endpoints(client):
    assert len(client.get("/api/reference/suppliers").json()) == 11
    assert len(client.get("/api/reference/programs").json()) == 25
    assert client.get("/api/reference/bom", params={"article_id": "P-00001046"}).json()
    assert client.get("/api/reference/links", params={"article_id": "P-00005775"}).json()[1]["supplier_id"] == "S-001033"
    assert client.get("/api/reference/plan", params={"program_id": "mass-00040633"}).json()
    assert client.post("/api/reference/refresh").json()["status"] == "ok"


def _legacy_pdp_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SOP - PDP"
    ws.append(["PF/SF", "S40-26", "S41-26", "2026-W42", "2027W01", "note"])
    ws.append(["Stators M3", 1000, 2000, 3000, 4000, "x"])
    ws.append(["Programme inconnu", 1, 2, 3, 4, ""])
    ws.append(["mass-00040922", 500, "", 700, 800, ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdp_import_and_activation(client):
    base = client.get("/api/articles/P-00043859/projection", params={"granularity": "week", "generate_proposals": "false"}).json()
    r = client.post("/api/pdp/import", files={"file": ("pdp.xlsx", _legacy_pdp_workbook())}, data={"name": "PDP S40"})
    assert r.status_code == 201, r.text
    rep = r.json()
    assert rep["created"] == 7 and rep["version"]["active"] and rep["version"]["programs"] == 2
    assert any("inconnu" in n for n in rep["notes"])
    after = client.get("/api/articles/P-00043859/projection", params={"granularity": "week", "generate_proposals": "false"}).json()
    d_base = dict(zip(base["periods"], next(s for s in base["series"] if s["key"] == "demand")["values"]))
    d_after = dict(zip(after["periods"], next(s for s in after["series"] if s["key"] == "demand")["values"]))
    assert d_after["2026-W40"] == pytest.approx(1000)          # qty_per = 1 for the stator stack
    assert d_after["2026-W40"] != d_base["2026-W40"]
    v = rep["version"]["id"]
    assert client.post(f"/api/pdp/versions/{v}/deactivate").json()["active"] is False
    back = client.get("/api/articles/P-00043859/projection", params={"granularity": "week", "generate_proposals": "false"}).json()
    assert next(s for s in back["series"] if s["key"] == "demand")["values"] == next(s for s in base["series"] if s["key"] == "demand")["values"]
    assert len(client.get(f"/api/pdp/versions/{v}/lines").json()) == 7
    assert client.delete(f"/api/pdp/versions/{v}").status_code == 204
    assert client.post("/api/pdp/import", files={"file": ("x.xlsx", b"not an excel")}).status_code == 422


def test_exports_and_reimport(client):
    r = client.get("/api/exports/simulation.xlsx", params={"planner": "QUENTIN", "granularity": "week",
                                                          "article_ids": ["P-00001046", "P-00005775"]})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) == {"PARAMETRES", "ARTICLES", "SIMULATION", "ALERTES", "CARNET_COMMANDES", "SAISIES"}
    ws = wb["SIMULATION"]
    assert ws["A4"].value == "P-00001046" and ws["C4"].value == "Besoin" and ws["C7"].value == "Commandes simulées"
    assert ws.cell(1, 5).value.startswith("2026-W")
    # the grid is made of formulas fed by the other sheets
    assert str(ws["E11"].value).startswith("=IF(PARAMETRES!$B$5")          # stock ferme
    assert "CARNET_COMMANDES" in str(ws["E5"].value) and "SAISIES" in str(ws["E8"].value)
    assert "OFFSET" in str(ws["E15"].value) and "COUNTIF" in str(ws["E16"].value)
    assert wb["ARTICLES"]["A2"].value == "P-00001046" and wb["ARTICLES"]["I2"].value > 0
    assert client.get("/api/exports/alerts.xlsx").status_code == 200
    assert client.get("/api/exports/orders.xlsx", params={"planner": "QUENTIN"}).status_code == 200
    # fill the SAISIES sheet, type simulated orders in the grid, post a receipt in the order book and re-import
    ws["F7"] = 1500          # P-00001046, second week → simulated order on the Monday of that week
    ws["G7"] = -250
    ws = wb["SAISIES"]
    ws.append(["COMMANDE", "P-00001046", "S-000545", "2026-10-20", 1600, "réimport", ""])
    ws.append(["RECEPTION", "P-00001046", "S-000545", "2026-09-22", 400, "", ""])
    ws.append(["AJUSTEMENT", "P-00003751", "", "2026-09-21", -30, "casse", ""])
    ws.append(["PRODUCTION", "mass-00040633", "", "2026-09-21", 250, "", ""])
    ws.append(["COMMANDE", "UNKNOWN", "", "2026-10-20", 5, "", ""])
    ws.append(["FOO", "P-00001046", "", "2026-10-20", 5, "", ""])
    wo = wb["CARNET_COMMANDES"]
    assert wo["A2"].value in ("P-00001046", "P-00005775")
    wo["J2"] = 250
    wo["K2"] = "2026-09-23"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/imports/entries", files={"file": ("simu.xlsx", buf.getvalue())})
    assert r.status_code == 201, r.text
    assert r.json()["created"] == 7 and r.json()["ignored"] == 2, r.json()
    orders = client.get("/api/entries/orders").json()
    assert len(orders) == 1 and orders[0]["source"] == "IMPORT"
    cells = client.get("/api/entries/cells", params={"article_id": "P-00001046"}).json()
    assert {(c["date"], c["qty"], c["source"]) for c in cells} == {("2026-09-21", 1500.0, "IMPORT"), ("2026-09-28", -250.0, "IMPORT")}
    receipts = client.get("/api/entries/receipts").json()
    assert len(receipts) == 2 and any(r["order_id"] == wo["D2"].value and r["qty"] == 250 for r in receipts)
    assert len(client.get("/api/entries/adjustments").json()) == 1
    assert len(client.get("/api/entries/production").json()) == 1
