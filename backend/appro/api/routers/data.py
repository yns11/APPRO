"""Paramètres › Données (synchronisation, orphans, calendar, adherence), supplier detail page and
generic Excel export of a screen table."""
from __future__ import annotations

import datetime as dt
import io
from typing import Any, Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from ...data.assembler import erp_dataset
from ...data.reference import HEADER_FILL, HEADER_FONT
from ...services import data_service, mrp_service
from ...services.context import AppContext
from .. import presenters as P
from .. import schemas as S
from ..deps import ctx_dep, session_dep
from ..fastjson import json_response

router = APIRouter(prefix="/api", tags=["data"])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------- Paramètres › Données
@router.get("/data/sync")
def sync(ctx: AppContext = Depends(ctx_dep), session: Session = Depends(session_dep)):
    """Age of every mirrored ERP table (``erp_sync_log``) ; ``late`` beyond two hours without a copy."""
    return data_service.sync_status(ctx, session)


@router.get("/data/orphans")
def orphans(ctx: AppContext = Depends(ctx_dep)):
    """Articles present in the ERP facts and absent from the reference table."""
    return data_service.orphan_articles(ctx)


@router.get("/data/calendar")
def calendar(ctx: AppContext = Depends(ctx_dep), session: Session = Depends(session_dep)):
    """Plant calendar : working weekdays and closure dates."""
    return data_service.plant_calendar(ctx, session)


@router.get("/data/adherence")
def adherence(weeks: int = Query(12, ge=1, le=52), bucket: Literal["week", "day"] = "week",
              plan_mode: Literal["monday", "latest"] = "monday", planner: str | None = None, program_id: str | None = None,
              article_id: str | None = None, ctx: AppContext = Depends(ctx_dep), session: Session = Depends(session_dep)):
    """Adherence of the reported consumption to the planned requirement (net, gross, bias, weekday
    profile and suggested ``demand_days``), per article, over the last complete ISO weeks."""
    return json_response(data_service.adherence(ctx, session, weeks=weeks, bucket=bucket, plan_mode=plan_mode,
                                                planner=planner, program_id=program_id, article_id=article_id))


# ---------------------------------------------------------------- supplier page
def _names(ctx: AppContext) -> dict[str, str]:
    df = ctx.table("ref_suppliers")
    return dict(zip(df["supplier_id"], df["name"]))


@router.get("/suppliers/{supplier_id}", response_model=S.SupplierDetail)
def supplier_detail(supplier_id: str, ctx: AppContext = Depends(ctx_dep), session: Session = Depends(session_dep)):
    """Everything about one supplier : record, linked articles with their sourcing rule and current
    situation, the supply flows of the day restricted to the supplier, the upcoming firm orders."""
    ds = erp_dataset(ctx.table)
    sup = next((s for s in ds.suppliers if s.supplier_id == supplier_id), None)
    if sup is None:
        raise HTTPException(404, f"Fournisseur inconnu : {supplier_id}")
    result = mrp_service.compute(ctx, session)
    names = _names(ctx)
    sups = {s.supplier_id: s for s in ds.suppliers}
    links = [l for l in ds.links if l.supplier_id == supplier_id]
    arts = {a.article_id: a for a in ds.articles}
    articles = []
    for l in links:
        a = arts.get(l.article_id)
        ar = result.articles.get(l.article_id)
        k = ar.kpis if ar else {}
        lane_plan = 0.0
        if ar:
            lane_plan = float(sum(float(ln.plan.sum() + ln.supply_proposed.sum()) for ln in ar.lanes if ln.supplier_id == supplier_id))
        articles.append(S.SupplierArticleRow(
            article_id=l.article_id, designation=a.designation if a else "", unit=a.unit if a else "", planner=a.planner if a else "",
            active=l.active, moq=l.moq, pack_qty=l.pack_qty, lead_time_days=l.lead_time(sups), lead_time_override=l.lead_time_days,
            quota_pct=l.quota_pct, priority=l.priority, max_qty_per_delivery=l.max_qty_per_delivery, max_qty_per_week=l.max_qty_per_week,
            supplier_count=sum(1 for x in ds.links if x.article_id == l.article_id and x.active),
            coverage_plan_days=k.get("coverage_plan_days"), stock_at_date=k.get("stock_at_date"), severity=k.get("severity"),
            first_stockout_plan=k.get("first_stockout_plan"), backlog_qty=k.get("backlog_qty", 0.0),
            open_firm_qty=float(sum(o.qty_open for o in (ar.lanes and [o for ln in ar.lanes if ln.supplier_id == supplier_id for o in ln.orders] or [])
                                    if o.order_type == "FIRM" and o.expected_date >= result.as_of and not o.ignored)) if ar else 0.0,
            planned_qty=lane_plan, proposal_count=sum(1 for p in (ar.proposals if ar else []) if p.supplier_id == supplier_id),
            urgent_count=sum(1 for p in (ar.proposals if ar else []) if p.supplier_id == supplier_id and p.urgent)))
    articles.sort(key=lambda r: ({"critical": 0, "warning": 1}.get(r.severity or "", 2), r.article_id))
    flows = P.flows_out(result, ds, names, None)

    def mine(rows):
        return [r for r in rows if (r.supplier_id or "") == supplier_id]

    upcoming = []
    for ar in result.articles.values():
        for ln in ar.lanes:
            if ln.supplier_id != supplier_id:
                continue
            for o in ln.orders:
                if o.order_type == "FIRM" and o.expected_date >= result.as_of:
                    upcoming.append(S.OrderRow(order_id=o.order_id, article_id=ar.article.article_id, designation=ar.article.designation,
                                               unit=ar.article.unit, planner=ar.article.planner, supplier_id=supplier_id,
                                               supplier_name=sup.name, order_type="FIRM", expected_date=o.expected_date,
                                               qty_ordered=o.qty_ordered, qty_open=o.qty_open, purch_id=o.ref))
    upcoming.sort(key=lambda o: (o.expected_date, o.article_id))
    return json_response(S.SupplierDetail(
        supplier=S.SupplierRef(supplier_id=sup.supplier_id, name=sup.name, country=sup.country, contact=sup.contact,
                               delivery_weekdays=",".join(str(d) for d in sorted(sup.delivery_weekdays)),
                               lead_time_days=sup.lead_time_days, active=sup.active),
        as_of=result.as_of, articles=articles, late=mine(flows.late), in_transit=mine(flows.in_transit),
        to_process=mine(flows.to_process), to_validate=mine(flows.to_validate), to_order=mine(flows.to_order),
        upcoming=upcoming[:200],
        kpis={"articles": len(articles), "active_links": sum(1 for l in links if l.active),
              "critical": sum(1 for r in articles if r.severity == "critical"),
              "late_orders": len(mine(flows.late)), "late_qty": float(sum(r.backlog for r in mine(flows.late))),
              "in_transit": len(mine(flows.in_transit)), "to_process": len(mine(flows.to_process)),
              "to_validate": len(mine(flows.to_validate)), "urgent_proposals": len(mine(flows.to_order)),
              "upcoming_orders": len(upcoming), "upcoming_qty": float(sum(o.qty_open for o in upcoming))}).model_dump(mode="json"))


# ---------------------------------------------------------------- generic Excel export of a screen table
class ExportColumn(BaseModel):
    key: str
    label: str
    type: Literal["text", "number", "date", "bool", "percent"] = "text"


class ExportTableIn(BaseModel):
    title: str = Field("Export", max_length=120)
    columns: list[ExportColumn] = Field(min_length=1, max_length=80)
    rows: list[dict[str, Any]] = Field(default_factory=list, max_length=50_000)
    footer: dict[str, Any] | None = None     # totals row (key → value), written after the rows in bold
    sheet: str = Field("Export", max_length=31)


def _cell_value(v: Any, kind: str) -> Any:
    if v is None or v == "":
        return None
    if kind == "bool":
        return "oui" if v in (True, "true", "oui", 1) else "non"
    if kind == "date":
        try:
            return dt.date.fromisoformat(str(v)[:10])
        except ValueError:
            return str(v)
    if kind in ("number", "percent"):
        try:
            return float(v)
        except (TypeError, ValueError):
            return str(v)
    return str(v) if not isinstance(v, (int, float)) else v


@router.post("/exports/table.xlsx")
def export_table(body: ExportTableIn):
    """Excel workbook of any screen table : the client sends the visible columns (in order) and the
    filtered rows ; the server types the cells (dates, numbers, percentages) and styles the sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = (body.sheet or "Export")[:31].replace("/", "-").replace(":", "-")
    for c, col in enumerate(body.columns, start=1):
        cell = ws.cell(1, c, col.label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(40, len(col.label) + 4))
    for r, row in enumerate(body.rows, start=2):
        for c, col in enumerate(body.columns, start=1):
            cell = ws.cell(r, c, _cell_value(row.get(col.key), col.type))
            if col.type == "date":
                cell.number_format = "dd/mm/yyyy"
            elif col.type == "number":
                cell.number_format = "#,##0.###"
            elif col.type == "percent":
                cell.number_format = "0.0"
    if body.footer:
        r = len(body.rows) + 2
        for c, col in enumerate(body.columns, start=1):
            v = body.footer.get(col.key)
            cell = ws.cell(r, c, _cell_value(v, col.type if col.type != "date" else "text") if v is not None else None)
            cell.font = Font(bold=True, color="1F2A44")
            if col.type == "number":
                cell.number_format = "#,##0.###"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(body.columns))}{max(1, len(body.rows) + 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    slug = "".join(ch if ch.isalnum() else "_" for ch in body.title.lower()).strip("_") or "export"
    fn = f"{slug}_{dt.date.today().isoformat()}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{fn}"'})
