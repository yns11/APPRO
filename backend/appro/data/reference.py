"""Reference tables managed in the application: frames, upsert / delete of rows, Excel templates
and imports.  One code path for the six tables, driven by :data:`appro.data.schemas.TABLES`."""
from __future__ import annotations

import datetime as dt
import io
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .schemas import REFERENCE_TABLES, TABLES, TableSchema, coerce
from .store import REF_MODELS, audit

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EXAMPLE_FONT = Font(italic=True, color="888888")


def schema_of(name: str) -> TableSchema:
    if name not in REFERENCE_TABLES:
        raise KeyError(name)
    return TABLES[name]


def frame(session: Session, name: str) -> pd.DataFrame:
    """The table as a canonical frame (coerced types)."""
    schema = schema_of(name)
    rows = session.scalars(select(REF_MODELS[name])).all()
    data = {c: [getattr(r, c) for r in rows] for c in schema.column_names}
    return coerce(pd.DataFrame(data, columns=list(schema.column_names)), schema)


def rows_out(session: Session, name: str) -> list[dict[str, Any]]:
    schema = schema_of(name)
    out = []
    for r in session.scalars(select(REF_MODELS[name])).all():
        d = {c: getattr(r, c) for c in schema.column_names}
        d["updated_by"], d["updated_at"] = r.updated_by, r.updated_at
        out.append(d)
    out.sort(key=lambda d: tuple(str(d[k]) for k in schema.key))
    return out


def _clean(schema: TableSchema, row: dict[str, Any]) -> dict[str, Any]:
    """Coerce one row to the canonical types ; missing keys raise."""
    df = coerce(pd.DataFrame([row]), schema)
    out = df.iloc[0].to_dict()
    for k in schema.key:
        if out[k] in ("", None):
            raise ValueError(f"clé manquante : {schema.column(k).label}")
    for c in schema.columns:
        if c.type == "date":
            v = out[c.name]
            out[c.name] = None if v is None or (isinstance(v, float) and pd.isna(v)) else v
        elif c.type in ("float", "int"):
            out[c.name] = float(out[c.name]) if c.type == "float" else int(out[c.name])
    return out


def upsert_row(session: Session, name: str, row: dict[str, Any], user: str) -> dict[str, Any]:
    """Create or update one row identified by the key columns.  Does not commit."""
    schema = schema_of(name)
    model = REF_MODELS[name]
    clean = _clean(schema, row)
    key = tuple(clean[k] for k in schema.key)
    obj = session.get(model, key if len(key) > 1 else key[0])
    action = "update" if obj is not None else "create"
    if obj is None:
        obj = model(**{k: clean[k] for k in schema.key})
        session.add(obj)
    for c in schema.column_names:
        setattr(obj, c, clean[c])
    obj.updated_by = user
    audit(session, user, action, name, "|".join(str(k) for k in key), clean.get("article_id"), clean)
    return clean


def delete_row(session: Session, name: str, key: dict[str, Any], user: str) -> bool:
    schema = schema_of(name)
    model = REF_MODELS[name]
    values = tuple(_clean(schema, {**{c: "" for c in schema.column_names}, **key})[k] for k in schema.key)
    obj = session.get(model, values if len(values) > 1 else values[0])
    if obj is None:
        return False
    audit(session, user, "delete", name, "|".join(str(v) for v in values), key.get("article_id"), key)
    session.delete(obj)
    return True


def replace_all(session: Session, name: str, df: pd.DataFrame, user: str) -> int:
    """Replace the whole table by the frame (already coerced).  Does not commit."""
    schema = schema_of(name)
    model = REF_MODELS[name]
    session.execute(delete(model))
    n = 0
    seen: set[tuple] = set()
    for row in df.to_dict("records"):
        clean = _clean(schema, row)
        key = tuple(clean[k] for k in schema.key)
        if key in seen:
            continue
        seen.add(key)
        session.add(model(**clean, updated_by=user))
        n += 1
    return n


def merge_all(session: Session, name: str, df: pd.DataFrame, user: str) -> int:
    n = 0
    for row in df.to_dict("records"):
        upsert_row(session, name, row, user)
        n += 1
    return n


# =============================================================================
# Excel
# =============================================================================
def _header_map(schema: TableSchema) -> dict[str, str]:
    """Accepted header spellings (canonical names and French labels, case-insensitive) → column."""
    m: dict[str, str] = {}
    for c in schema.columns:
        m[c.name.lower()] = c.name
        m[c.label.lower()] = c.name
    return m


def template_workbook(name: str, rows: list[dict[str, Any]] | None = None) -> bytes:
    """Minimal template of one reference table: French headers (one per canonical column), one
    example row (or the current content), plus a NOTICE sheet describing every column."""
    schema = schema_of(name)
    wb = Workbook()
    ws = wb.active
    ws.title = schema.label[:31]
    for c, col in enumerate(schema.columns, start=1):
        cell = ws.cell(1, c, col.label)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(34, len(col.label) + 4))
    if rows:
        for r, row in enumerate(rows, start=2):
            for c, col in enumerate(schema.columns, start=1):
                v = row.get(col.name)
                ws.cell(r, c, "oui" if v is True else "non" if v is False else v)
    else:
        for c, col in enumerate(schema.columns, start=1):
            v = col.example
            ws.cell(2, c, "oui" if v is True else "non" if v is False else v).font = EXAMPLE_FONT
    for c, col in enumerate(schema.columns, start=1):
        if col.type == "date":
            for r in range(2, 400):
                ws.cell(r, c).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    notice = wb.create_sheet("NOTICE")
    notice.append([schema.label, schema.description])
    notice.append([])
    notice.append(["Colonne", "Nom technique", "Type", "Obligatoire", "Description"])
    for c in range(1, 6):
        notice.cell(3, c).font = Font(bold=True)
    for col in schema.columns:
        notice.append([col.label, col.name, {"str": "texte", "float": "nombre", "int": "entier", "bool": "oui / non",
                                             "date": "date (AAAA-MM-JJ)"}[col.type],
                       "oui" if col.required or col.name in schema.key else "non", col.description])
    notice.append([])
    notice.append(["Clé", ", ".join(schema.column(k).label for k in schema.key),
                   "Une ligne par clé ; l'import « remplacer » vide la table puis la recharge, « fusionner » met à jour les clés présentes."])
    notice.append(["Exemple", "", "La ligne 2 en italique est un exemple à remplacer (elle est ignorée si elle reste identique)."])
    for c, w in enumerate((26, 24, 16, 12, 80), start=1):
        notice.column_dimensions[get_column_letter(c)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workbook(name: str, content: bytes) -> tuple[pd.DataFrame, list[str]]:
    """Read the first sheet of an import file: headers = French labels or canonical names."""
    schema = schema_of(name)
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    notes: list[str] = []
    if not rows:
        return pd.DataFrame(columns=list(schema.column_names)), ["classeur vide"]
    hmap = _header_map(schema)
    cols: dict[int, str] = {}
    for i, h in enumerate(rows[0]):
        key = str(h).strip().lower() if h is not None else ""
        if key in hmap:
            cols[i] = hmap[key]
        elif key:
            notes.append(f"colonne ignorée : {h!r}")
    missing = [c.label for c in schema.columns if (c.required or c.name in schema.key) and c.name not in cols.values()]
    if missing:
        raise ValueError(f"colonnes obligatoires absentes : {', '.join(missing)}")
    example = {c.name: c.example for c in schema.columns}
    data = []
    for n, r in enumerate(rows[1:], start=2):
        if r is None or all(v in (None, "") for v in r):
            continue
        row = {col: (r[i] if i < len(r) else None) for i, col in cols.items()}
        if n == 2 and all(str(row.get(k, "")) == str(example.get(k, "")) for k in schema.key):
            notes.append("ligne d'exemple ignorée")
            continue
        if any(row.get(k) in (None, "") for k in schema.key):
            notes.append(f"ligne {n} : clé manquante, ignorée")
            continue
        data.append({k: (v.date() if isinstance(v, dt.datetime) else v) for k, v in row.items()})
    df = coerce(pd.DataFrame(data, columns=list(cols.values())), schema)
    return df, notes
